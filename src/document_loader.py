"""Extraction of PDF, DOCX, and plain-text uploads with page metadata."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document


def extract_document(uploaded_file) -> tuple[list[dict], list[str]]:
    """Return non-empty page records and user-safe warnings for one upload."""
    name = uploaded_file.name
    suffix = Path(name).suffix.lower()
    warnings: list[str] = []
    pages: list[dict] = []
    try:
        raw = uploaded_file.getvalue()
        if suffix == ".pdf":
            with fitz.open(stream=raw, filetype="pdf") as pdf:
                for number, page in enumerate(pdf, start=1):
                    text = page.get_text("text").strip()
                    if text:
                        pages.append({"file_name": name, "page_number": number, "text": text})
        elif suffix == ".docx":
            document = Document(BytesIO(raw))
            text = "\n".join(p.text for p in document.paragraphs if p.text.strip()).strip()
            for table in document.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text.strip() for cell in row.cells)
            if text.strip():
                pages.append({"file_name": name, "page_number": 1, "text": text.strip()})
        elif suffix == ".txt":
            text = raw.decode("utf-8", errors="replace").strip()
            if text:
                pages.append({"file_name": name, "page_number": 1, "text": text})
        else:
            warnings.append(f"{name}: unsupported file type. Use PDF, DOCX, or TXT.")
    except Exception as exc:
        warnings.append(f"{name}: could not be read ({exc.__class__.__name__}).")
    if not pages and not warnings:
        warnings.append(f"{name}: no extractable text was found.")
    return pages, warnings


def extract_documents(uploaded_files) -> tuple[list[dict], list[str]]:
    pages, warnings = [], []
    for uploaded_file in uploaded_files or []:
        file_pages, file_warnings = extract_document(uploaded_file)
        pages.extend(file_pages)
        warnings.extend(file_warnings)
    return pages, warnings
